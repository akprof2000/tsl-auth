"""Сборка Word-версии ЧТЗ из Markdown: по умолчанию docs/TZ-tsl-auth.docx из docs/tz.md.

Запуск из корня репозитория: python scripts/build-tz-docx.py [--version 2.2] [--date 30.09.2026]
Другое ЧТЗ: python scripts/build-tz-docx.py --src docs/task-active-directory.md --out docs/TZ-active-directory.docx     --version 1.0 --status "постановка, реализация не начата"
Нужен python-docx (pip install python-docx). Переводит заголовки, абзацы, списки, таблицы GFM, блоки кода и картинки
(схемы Mermaid уже отрисованы в docs/diagrams — см. scripts/render-diagrams.ps1; исходники схем в <details> опускаются).
Титульный блок (название, версия, статус) добавляется здесь; ссылки превращаются в текст.
"""
import argparse
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "docs", "tz.md")
OUT = os.path.join(ROOT, "docs", "TZ-tsl-auth.docx")
TASK = re.compile(r"^\[( |x|X)\]\s+")

LINK = re.compile(r"\[([^\]]+)\]\([^)]+\)")
IMAGE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def inline(paragraph, text):
    """Жирный (**…**) и код (`…`) внутри абзаца; ссылки — только текст."""
    text = LINK.sub(r"\1", text)
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def table_rows(lines):
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def build(version, date, src=SRC, out=OUT, status="проект для согласования"):
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    with open(src, encoding="utf-8") as f:
        lines = f.read().split("\n")

    title = lines[0].lstrip("# ").strip()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Техническое задание").bold = True
    doc.add_paragraph(title).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("ERP «Уклад»: компонент «Управление идентификацией»").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Версия {version} от {date}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Статус: {status}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Составлено по шаблону ЧТЗ базы знаний UKA (07.01).").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    figure = 0
    i = 1
    in_details = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("<details"):
            in_details = True
        if in_details:
            if stripped.startswith("</details>"):
                in_details = False
            i += 1
            continue

        if stripped.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            i += 1
            continue

        m = IMAGE.match(stripped)
        if m:
            alt, path = m.groups()
            full = os.path.join(os.path.dirname(src), path)
            if os.path.exists(full):
                figure += 1
                doc.add_picture(full, width=Cm(16.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph(f"Рисунок {figure}. {alt}")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            doc.add_heading(stripped.lstrip("#").strip(), level=min(level - 1, 3) if level > 1 else 1)
            i += 1
            continue

        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = table_rows(block)
            if rows:
                table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                table.style = "Table Grid"
                for r, cells in enumerate(rows):
                    for c, text in enumerate(cells):
                        cell = table.cell(r, c)
                        cell.text = ""
                        inline(cell.paragraphs[0], text)
                        if r == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
                doc.add_paragraph()
            continue

        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            indent, marker, text = m.groups()
            style_name = "List Number" if marker[0].isdigit() else "List Bullet"
            if indent:
                style_name += " 2"
            p = doc.add_paragraph(style=style_name)
            # Пункт-галочка «- [ ]» / «- [x]» — знаком флажка.
            text = TASK.sub(lambda t: "☑ " if t.group(1) != " " else "☐ ", text)
            inline(p, text)
            i += 1
            continue

        if stripped == "":
            i += 1
            continue

        # Абзац: соседние строки склеиваются до пустой строки.
        para = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(\s*([-*]|\d+\.)\s+|#|\||!\[|```|<details)", lines[i]):
            para.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        inline(p, " ".join(para))

    doc.save(out)
    print(f"{out}: {figure} рисунков, версия {version} от {date}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--version", default="2.2")
    parser.add_argument("--date", default="30.09.2026")
    parser.add_argument("--src", default=SRC, help="исходный Markdown")
    parser.add_argument("--out", default=OUT, help="файл .docx")
    parser.add_argument("--status", default="проект для согласования")
    args = parser.parse_args()
    sys.exit(build(args.version, args.date, os.path.abspath(args.src), os.path.abspath(args.out), args.status))
